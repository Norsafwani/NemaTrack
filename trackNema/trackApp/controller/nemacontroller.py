# Create your views here. where all the functions lies
from django.http.response import HttpResponse, JsonResponse
from django.shortcuts import render, redirect
from datetime import date, datetime
from trackApp.models import AuthUser, AuthPermission
from django.contrib.auth.forms import AuthenticationForm 
from django.contrib.auth import logout, authenticate, login
from django.core.files.storage import FileSystemStorage
from trackApp.models import Nema, Returnformnema
import openpyxl #For Upload Excel
from django.contrib import messages

# from trackApp.forms import DocumentForm
# from trackApp.forms import DocumentForm

#Function for Display Nema Data
def indexnema(request):
    nema = Nema.objects.all()
    # print('print this', product)
    content = {
        'nema':nema,
    }
    return render(request, 'nema/nema_data.html', content)

#Function Date
# def strdatetimeToDatetime(strdate):
#     datetime2 = datetime.strptime(strdate, '%Y-%m-%d')
#     # return datetime2

#Function for html create new Nema
def nemacreate(request):
    return render(request,'nema/nema_create.html',{})

#Function Submit Create New Nema
def submitnema(request):
    # id = request.session['id']
    # if request.session._session:
        if request.method=='POST':
            # nema_id = request.POST['id']
            devui_no = request.POST['devui']
            appkey = request.POST['app_key']
            shipdatereceived = request.POST['ship_date_received']
            siteinstalldate = request.POST['site_install_date']
            datedeliver = request.POST['date_deliver']           
            lightsolname = request.POST['lightsol_name']           
            licenseactivedate = request.POST['license_active_date']          
            licenseexpireddate = request.POST['license_expired_date']          
            contractorname = request.POST['contractor_name']           
            endclientname = request.POST['end_client_name']           
            projecttendername = request.POST['project_tender_name']           
            donumber = request.POST['do_number']            
            remarks = request.POST['remarks']

            # For insert into database-id= nema_id,
            # object = nama model( namacolumn=nama variable, others - if any)
            track = Nema(devui=devui_no, app_key=appkey, ship_date_received=shipdatereceived, site_install_date= siteinstalldate, 
            date_deliver=datedeliver, lightsol_name = lightsolname, license_active_date=licenseactivedate, license_expired_date= licenseexpireddate,
            contractor_name= contractorname,  end_client_name= endclientname, project_tender_name= projecttendername, 
            do_number= donumber, remarks= remarks )
            track.save()
            return redirect('/indexnema')


#Function for Delete Nema Data
def deletenema(request, nema_id):
    deleteN = Nema.objects.get(nema_id=nema_id)
    deleteN.delete()
    # success_url = reverse_lazy('indexproduct')
    return redirect('/indexnema')

#Function for Update Nema Data
def updatenema(request,nema_id):
    nema = Nema.objects.get(nema_id=nema_id)
   
    obj = {
        
        'nema': nema,
        
    }
    return render(request, 'nema/nema_update.html', obj)

#Funtion for Submit Update Nema Data
def updatesubmitnema(request,id):
    if request.method == "POST":
        nema_id = request.POST['nema_id']
        devui_no = request.POST['devui']       
        appkey = request.POST['app_key']       
        shipdatereceived = request.POST['ship_date_received']
        siteinstalldate = request.POST['site_install_date']
        datedeliver = request.POST['date_deliver']         
        lightsolname = request.POST['lightsol_name']      
        licenseactivedate = request.POST['license_active_date']      
        licenseexpireddate = request.POST['license_expired_date']       
        contractorname = request.POST['contractor_name']      
        endclientname = request.POST['end_client_name']   
        projecttendername = request.POST['project_tender_name']     
        donumber = request.POST['do_number']    
        remarks = request.POST['remarks']
             
        # print(devui)
        nema = Nema.objects.get(nema_id=nema_id)
        nema.devui = devui_no
        nema.app_key = appkey
        nema.ship_date_received = shipdatereceived
        nema.site_install_date = siteinstalldate
        nema.date_deliver = datedeliver
        nema.lightsol_name = lightsolname
        nema.license_active_date = licenseactivedate
        nema.license_expired_date = licenseexpireddate
        nema.contractor_name = contractorname
        nema.end_client_name = endclientname
        nema.project_tender_name = projecttendername
        nema.do_number = donumber
        nema.remarks = remarks
        # nema.date_modified = now
        # nema.modified_by = user_id
        nema.save()
    return redirect('/indexnema')


#Function for Searching 
def searchnema(request):
    # global nemas
    if request.method == "POST":
        searchNema =  request.POST['searchNema'] 
        #nemas(Return search result)
        nemas = Nema.objects.filter(devui__contains=searchNema)
        nemas2 = Nema.objects.filter(app_key__contains=searchNema)

        # app_key__contains=searchNema
        obj = {
        'searchNema': searchNema,
        'nemas': nemas,
        'nemas2': nemas2,
        }

        return render(request, 'nema/nema_search.html', obj)
    else:
        return render(request, 'nema/nema_search.html', {})

#Function for html create new Nema
# def nemacreate(request):
#     return render(request,'nema/nema_create.html',{})

#Function fo view 1 specific Nema
def viewnema(request,nema_id):
    # admin = ''
    objnema = Nema.objects.get(nema_id=nema_id)
    obj = {
        'objnema': objnema
    }
    return render(request, 'nema/nema_view.html', obj)

#Function 'Return Form Nema'
def returnformnema(request,nema_id):
    # admin = ''
    objNema = Nema.objects.get(nema_id=nema_id)
    obj = {
        'objNema': objNema
         }
    return render(request, 'nema/nema_form.html', obj)

 #Function for Save Return Form info
def submitreturnform(request):
    if request.method=='POST':
        date_uninstall = request.POST['dateuninstall']
        print(date_uninstall)

        date_detect = request.POST['datedetect']   
        print(date_detect)

        proofdescribe = request.POST['proof_describe']          
        nosiri = request.POST['no_siri']
        image_proof = request.FILES['image_proof']

        # masuk ke database
        # object = nama model( namacolumn=nama variable, others - if any)
        form = Returnformnema(dateuninstall=date_uninstall, datedetect=date_detect, proof_describe=proofdescribe,
                                  no_siri=nosiri, image_proof=image_proof)
        form.save()
        return redirect('return_nema')

#Function for upload File (File System Storage) error=( MultiValueDictKeyError at /submitreturnform - image_proof)
# def fileupload(request):
#     if request.method == 'POST' and request.FILES['image_proof']:
#         image_proof = request.FILES['image_proof']
#         fs = FileSystemStorage()
#         filename = fs.save(image_proof.name, image_proof)
#         uploaded_file_url = fs.url(filename)

#         return render(request, 'nema_form.html', {
#             'uploaded_file_url': uploaded_file_url
#         })
#     return render(request, 'nema/nema_form.html')


#EXAMPLE FORM - SUBMIT RETURN FORM
def submitreturnform(request):

    date_uninstall = request.POST['dateuninstall']
    print(date_uninstall)
    date_detect = request.POST['datedetect']   
    print(date_detect)
    proofdescribe = request.POST['proof_describe']          
    nosiri = request.POST['no_siri']
    img_exist = request.FILES.get('image_proof','')
    try:

        filename = None
        if img_exist != '':
            the_file = request.FILES['image_proof']
            fs = FileSystemStorage()
            path = 'trackNema/static/images/'
            filename = fs.save(path+username+'-'+the_file.name, the_file)
            uploaded_file_url = fs.url(filename)
        
        formsave = Returnformnema(dateuninstall=date_uninstall, datedetect=date_detect, proof_describe=proofdescribe,
                                  no_siri=nosiri, img=filename)
        formsave.save()
        # messages.success(request, 'Successfully submit report')
        return redirect('/return_nema')
    except Exception as e:
        # return HttpResponse(e)
        # messages.error(request, 'Error submitting report with error code '+str(e))
        print("ERROR")
        return redirect('/return_nema')

#Function for Display Return Nema [List]
def return_nema(request):
    returnnema = Returnformnema.objects.all()
    # print('print this', product)
    content = {
        'returnnema':returnnema,
    }
    return render(request, 'nema/return_form_list.html', content)

#Function for upload File (File System Storage)
# def fileupload(request):
#     if request.method == 'POST' and request.FILES['documents']:
#         documents = request.FILES['documents']
#         fs = FileSystemStorage()
#         filename = fs.save(documents.name, documents)
#         uploaded_file_url = fs.url(filename)

#         return render(request, 'nema_form.html', {
#             'uploaded_file_url': uploaded_file_url
#         })
#     return render(request, 'nema/nema_form.html')


# #TRY-Function for Upload Files & Images ni ke .. kn ?
# def form_upload(request):
#     if request.method == 'POST':
#         form = TrackappDocument(request.POST, request.FILES)
#         if form.is_valid():
#            form.save()
#            return redirect('indexnema')
#     else:
#         form = TrackappDocument()
#     return render(request, 'nema/nema_form.html', {
#         'form': form
#     })

#Function for upload File 
# def form_upload(request):
#     if request.method == 'POST' and request.FILES['file']:
#         file = request.FILES['file']
#         fs = FileSystemStorage()
#         filename = fs.save(file.name, file)
#         uploaded_file_url = fs.url(filename)

#         # ni klau berjaya ke? x ingt dh aku home la ei
#         return render(request, 'home.html', {  
#             'uploaded_file_url': uploaded_file_url
#         })

#     #  yang ni?   html dia lain sikit kn
#     return render(request, 'home.html')

#TRY UPLOAD - MODEL - Document
#Function for Display Return Nema [List]
# def return_nema(request):
#     returnnema = Document.objects.all()
#     # print('print this', product)
#     content = {
#         'returnnema':returnnema,
#     }
#     return render(request, 'nema/return_form_list.html', content)

# def upload_return(request):
#     if request.method == 'POST':
#         form = DocumentForm(request.POST, request.FILES)
#         if form.is_valid():
#             form.save()
#             return redirect('/nema')
#     else:
#         form = DocumentForm()
#     return render(request, 'nema/upload_try.html', {'form': form })

    
# TRY- Function for upload Excel(1) [ sepang_iot] --Error = BadZipFile at /get_excel
# @login_required 
def get_excel(request):
    import openpyxl

  #(adddevicetype?)
    adddevicetype = request.POST['adddevicetype2']
    # return HttpResponse(adddevicetype)

    excel_file = request.FILES["excel_file"]
    wb = openpyxl.load_workbook(excel_file)
    worksheet = wb["Sheet1"] #Sheet1 = empty file excel tab = sheet1 
    count = countsuccess = countfail = 0 
    msg = 'none'

    for row in worksheet.iter_rows(min_row=2,values_only=True):
        count = count + 1
        devui = str(row[0])
        app_key = str(row[1])
        ship_date_received = str(row[2])
        site_install_date = str(row[3])
        date_deliver = str(row[4])
        lightsol_name = str(row[5])
        license_active_date = str(row[6])
        license_expired_date = str(row[7])
        contractor_name = str(row[8])
        end_client_name = str(row[9])
        project_tender_name = str(row[10])
        do_number = str(row[11])
        remarks = str(row[12])

        nema = Nema(devui=devui,app_key=app_key,
                        ship_date_received=ship_date_received,
                        site_install_date=site_install_date,date_deliver=date_deliver,
                        lightsol_name=lightsol_name,license_active_date=license_active_date,
                        license_expired_date=license_expired_date,
                        contractor_name=contractor_name,end_client_name=end_client_name,
                        project_tender_name=project_tender_name, do_number=do_number,remarks=remarks)
        nema.save()
        msg = 'Successfully add devices'
        countsuccess += 1
        
    # return HttpResponse(payload)
    # return HttpResponseRedirect('/')
    msg = 'Success add: '+str(countsuccess)+' - Failed add: '+str(countfail)
    func.logaction(userid, 'createdeviceexcel', countsuccess)
    return HttpResponse(msg)


#TRY- Function File Upload
# def my_view(request):
#     # print(f"Great! You're using Python 3.6+. If you fail here, use the right version.")
#     message = 'Upload as many files as you want!'
#     # Handle file upload
#     if request.method == 'POST':
#         form = NemareturnformForm(request.POST, request.FILES)
#         if form.is_valid():
#             newdoc = Nemareturnform(docfile=request.FILES['docfile'])
#             newdoc.save()

#             # Redirect to the document list after POST
#             return redirect('my-view')
#         else:
#             message = 'The form is not valid. Fix the following error:'
#     else:
#         form = NemareturnformForm()  # An empty, unbound form

#     # Load documents for the list page
#     documents = Nemareturnform.objects.all()

#     # Render list page with the documents and the form
#     context = {'documents': documents, 'form': form, 'message': message}
#     return render(request, 'return_form_list.html', context)



